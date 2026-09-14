# 엑셀(XLSX) 산출물 생성
#   pip install openpyxl
#   python 이_파일.py
#
# '문서' 시트에 본문이 흐르고, 표는 각각 별도 시트로 떨어진다.
__PARSER__

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

INK, DIM, ACC, BAND = "FF1C1B19", "FF6B6862", "FF2383E2", "FFFAFAF8"
FONT = "맑은 고딕"
thin = Side(style="thin", color="FFE4E2DD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

meta, blocks = parse(SRC)
wb = Workbook()
ws = wb.active
ws.title = "문서"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 110

r = 2
def put(text, size=11, bold=False, color=INK, indent=0, height=None, wrap=True):
    global r
    c = ws.cell(row=r, column=2, value=text)
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=wrap, vertical="top", indent=indent)
    if height:
        ws.row_dimensions[r].height = height
    r += 1

put(TITLE, 20, True, INK, height=30)
bits = [meta.get(k, "") for k in ("id", "phase", "status", "owner", "updated")]
put("   ·   ".join([b for b in bits if b]), 9, False, DIM, height=18)
r += 1

tables = []
for kind, val in blocks:
    if kind == "h1":
        continue
    elif kind == "h2":
        r += 1
        put(val, 14, True, INK, height=24)
    elif kind == "h3":
        put(val, 12, True, INK, height=20)
    elif kind == "h4":
        put(val, 11, True, DIM, height=18)
    elif kind == "p":
        put(val, 10, height=max(15, 15 * (len(val) // 90 + 1)))
    elif kind == "quote":
        put("│ " + val, 10, False, DIM)
    elif kind == "list":
        for it in val:
            put("•  " + it, 10, indent=1)
    elif kind == "code":
        for line in val.split("\n"):
            c = ws.cell(row=r, column=2, value=line)
            c.font = Font(name="Consolas", size=9, color=INK)
            c.fill = PatternFill("solid", fgColor=BAND)
            r += 1
    elif kind == "table" and val:
        tables.append(val)
        put("→ 표는 '표%d' 시트에 있습니다" % len(tables), 9, False, ACC)

# ── 표: 시트 하나에 하나씩. 엑셀 표 기능(정렬·필터)을 켜 둔다 ─────────
for n, rows in enumerate(tables, 1):
    s = wb.create_sheet("표%d" % n)
    s.sheet_view.showGridLines = False
    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            c = s.cell(row=ri, column=ci, value=val)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if ri == 1:
                c.font = Font(name=FONT, size=10, bold=True, color="FFFFFFFF")
                c.fill = PatternFill("solid", fgColor=INK)
            else:
                c.font = Font(name=FONT, size=10, color=INK)
    width = len(rows[0])
    for ci in range(1, width + 1):
        longest = max((len(str(row[ci - 1])) for row in rows if ci - 1 < len(row)), default=10)
        s.column_dimensions[get_column_letter(ci)].width = min(52, max(12, longest + 4))
    ref = "A1:%s%d" % (get_column_letter(width), len(rows))
    if len(rows) > 1:
        t = Table(displayName="T%d" % n, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        s.add_table(t)
    s.freeze_panes = "A2"

wb.save(OUT + ".xlsx")
print("만들었습니다: " + OUT + ".xlsx  (표 %d개)" % len(tables))
