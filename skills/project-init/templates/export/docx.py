# 워드(DOCX) 산출물 생성
#   pip install python-docx
#   python 이_파일.py
__PARSER__

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK, DIM, ACC = RGBColor(0x1C, 0x1B, 0x19), RGBColor(0x6B, 0x68, 0x62), RGBColor(0x23, 0x83, 0xE2)
FONT = "맑은 고딕"

meta, blocks = parse(SRC)
doc = Document()

for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.2)
    s.left_margin = s.right_margin = Cm(2.4)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(10)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)

def para(text, size=10, bold=False, color=INK, before=0, after=6, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if indent:
        p.paragraph_format.left_indent = Cm(0.6 * indent)
    r = p.add_run(text)
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, FONT
    r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return p

def shade(cell, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)

para(TITLE, 22, True, INK, after=2)
bits = [meta.get(k, "") for k in ("id", "phase", "status", "owner", "updated")]
para("   ·   ".join([b for b in bits if b]), 9, False, DIM, after=16)

for kind, val in blocks:
    if kind == "h1":
        continue
    elif kind == "h2":
        para(val, 15, True, INK, before=16, after=4)
    elif kind == "h3":
        para(val, 12, True, INK, before=10, after=3)
    elif kind == "h4":
        para(val, 10.5, True, DIM, before=8, after=2)
    elif kind == "p":
        para(val)
    elif kind == "quote":
        para("│  " + val, 10, False, DIM, indent=1)
    elif kind == "list":
        for it in val:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(it)
            r.font.size, r.font.name = Pt(10), FONT
            r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    elif kind == "code":
        for line in val.split("\n"):
            p = para(line, 9, color=INK, after=0)
            p.runs[0].font.name = "Consolas"
    elif kind == "table" and val:
        t = doc.add_table(rows=len(val), cols=len(val[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for ri, row in enumerate(val):
            for ci in range(len(val[0])):
                cell = t.cell(ri, ci)
                cell.text = row[ci] if ci < len(row) else ""
                shade(cell, "1C1B19" if ri == 0 else ("FAFAF8" if ri % 2 == 0 else "FFFFFF"))
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    for r in p.runs:
                        r.font.size, r.font.name = Pt(9), FONT
                        r.font.bold = ri == 0
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else INK
                        r.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        doc.add_paragraph()

doc.save(OUT + ".docx")
print("만들었습니다: " + OUT + ".docx")
