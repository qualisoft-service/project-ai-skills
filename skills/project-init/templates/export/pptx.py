# 발표자료(PPTX) 산출물 생성
#   pip install python-pptx
#   python 이_파일.py
#
# H2 하나가 슬라이드 하나가 된다. 표는 슬라이드 표로, 목록은 불릿으로 들어간다.
__PARSER__

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

W, H = Inches(13.333), Inches(7.5)
INK   = RGBColor(0x1C, 0x1B, 0x19)
DIM   = RGBColor(0x6B, 0x68, 0x62)
ACC   = RGBColor(0x23, 0x83, 0xE2)
BAND  = RGBColor(0xFA, 0xFA, 0xF8)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
FONT  = "Malgun Gothic"

meta, blocks = parse(SRC)
prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]

def textbox(slide, x, y, w, h, text, size, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = color
    return tf

def bar(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s

# ── 표지 ─────────────────────────────────────────────────────────────
cover = prs.slides.add_slide(BLANK)
bar(cover, 0, 0, Inches(0.18), H, ACC)
textbox(cover, Inches(1.1), Inches(2.5), Inches(11), Inches(1.6), TITLE, 40, INK, True)
bits = [meta.get(k, "") for k in ("id", "phase", "owner", "updated")]
textbox(cover, Inches(1.15), Inches(4.1), Inches(11), Inches(0.6),
        "   ·   ".join([b for b in bits if b]), 13, DIM)
if meta.get("summary"):
    textbox(cover, Inches(1.15), Inches(4.7), Inches(10.5), Inches(1.2), meta["summary"], 15, INK)

# ── 본문: H2 단위로 슬라이드 ──────────────────────────────────────────
def new_slide(title):
    s = prs.slides.add_slide(BLANK)
    bar(s, Inches(0.9), Inches(0.62), Inches(0.055), Inches(0.42), ACC)
    textbox(s, Inches(1.1), Inches(0.5), Inches(11.4), Inches(0.7), title, 24, INK, True)
    return s, Inches(1.45)

def add_table(slide, rows, top):
    cols = len(rows[0])
    shape = slide.shapes.add_table(len(rows), cols, Inches(1.1), top,
                                   Inches(11.2), Inches(0.36) * len(rows))
    tbl = shape.table
    for c in range(cols):
        cell = tbl.cell(0, c)
        cell.text = rows[0][c]
        cell.fill.solid(); cell.fill.fore_color.rgb = INK
        runs = cell.text_frame.paragraphs[0].runs   # 빈 셀은 run 이 없다
        if runs:
            runs[0].font.size, runs[0].font.bold, runs[0].font.name = Pt(11), True, FONT
            runs[0].font.color.rgb = PAPER
    for r in range(1, len(rows)):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = rows[r][c] if c < len(rows[r]) else ""
            cell.fill.solid()
            cell.fill.fore_color.rgb = BAND if r % 2 == 0 else PAPER
            if cell.text_frame.paragraphs[0].runs:
                pr = cell.text_frame.paragraphs[0].runs[0]
                pr.font.size, pr.font.name = Pt(10), FONT
                pr.font.color.rgb = INK
    return top + Inches(0.36) * len(rows) + Inches(0.3)

slide, y = None, Inches(1.45)
for kind, val in blocks:
    if kind == "h1":
        continue
    if kind == "h2":
        slide, y = new_slide(val)
        continue
    if slide is None:
        slide, y = new_slide("개요")
    if y > Inches(6.7):                      # 넘치면 이어지는 슬라이드로
        slide, y = new_slide("(이어서)")
    if kind in ("h3", "h4"):
        textbox(slide, Inches(1.1), y, Inches(11.2), Inches(0.4), val, 15, INK, True)
        y += Inches(0.45)
    elif kind == "p":
        lines = max(1, len(val) // 70 + 1)
        textbox(slide, Inches(1.1), y, Inches(11.2), Inches(0.3) * lines, val, 12, INK)
        y += Inches(0.3) * lines + Inches(0.12)
    elif kind == "quote":
        bar(slide, Inches(1.1), y, Inches(0.04), Inches(0.34), ACC)
        textbox(slide, Inches(1.3), y, Inches(11), Inches(0.34), val, 12, DIM)
        y += Inches(0.5)
    elif kind == "list":
        for it in val:
            textbox(slide, Inches(1.25), y, Inches(11), Inches(0.3), "•  " + it, 12, INK)
            y += Inches(0.32)
        y += Inches(0.14)
    elif kind == "code":
        textbox(slide, Inches(1.1), y, Inches(11.2), Inches(0.3), val.split("\n")[0], 11, DIM)
        y += Inches(0.4)
    elif kind == "table" and val:
        y = add_table(slide, val[:14], y)

prs.save(OUT + ".pptx")
print("만들었습니다: " + OUT + ".pptx")
